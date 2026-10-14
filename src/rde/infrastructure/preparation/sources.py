"""Lexical source loading, key candidates and stable record references."""

import csv
import hashlib
import io
import itertools
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from rde.domain.models.preparation import Source

MAX_ROWS = 250_000
MAX_CELLS = 5_000_000


def canonical(value):
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    )


def digest(value):
    return hashlib.sha256(canonical(value).encode()).hexdigest()


def file_hash(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def empty(value):
    return value is None or (isinstance(value, str) and (value == "" or value.isspace()))


def normalize(value, policy="exact"):
    if empty(value):
        return None
    text = str(value)
    if policy != "exact":
        text = text.strip()
    if policy == "nfc_trim":
        text = unicodedata.normalize("NFC", text)
    return text


def key(values, columns, policy="exact"):
    result = tuple(normalize(values[c], policy) for c in columns)
    return None if None in result else result


@dataclass
class Row:
    index: int
    values: dict
    sha256: str


@dataclass
class Table:
    source: Source
    columns: list[str]
    rows: list[Row]
    numeric_cells: dict[str, int]

    def ref(self, row, column=None):
        return {
            "source_id": self.source.id,
            "source_sha256": self.source.sha256,
            "sheet": self.source.sheet,
            "record_index": row.index,
            "row_sha256": row.sha256,
            **({"column": column} if column is not None else {}),
        }

    def require_columns(self, columns):
        missing = set(columns) - set(self.columns)
        if missing:
            raise ValueError(f"source {self.source.id} lacks columns: {sorted(missing)}")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("source contains non-finite numeric values")
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def load(source: Source, snapshot: Path | None = None):
    path = snapshot or Path(source.path)
    if not path.is_absolute() or not path.is_file():
        raise ValueError("source must be an existing absolute file")
    raw = path.read_bytes()
    if len(raw) != source.bytes or hashlib.sha256(raw).hexdigest() != source.sha256:
        raise ValueError(f"SOURCE_CHANGED: {source.id}; create a new source revision")
    suffix = Path(source.path).suffix.lower()
    numeric = Counter()
    records = []
    if suffix in {".csv", ".tsv"}:
        reader = csv.reader(
            io.StringIO(raw.decode(source.encoding), newline=""),
            delimiter=source.delimiter,
            strict=True,
        )
        try:
            columns = next(reader)
        except StopIteration:
            raise ValueError("source has no header") from None
        for row in reader:
            if len(row) != len(columns):
                raise ValueError("CSV record width differs from header; no silent skipped records")
            records.append(row)
            if len(records) > MAX_ROWS or len(records) * len(columns) > MAX_CELLS:
                raise ValueError("source exceeds preparation row/cell limit")
    elif suffix == ".xlsx":
        import openpyxl

        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            if sum(item.file_size for item in package.infolist()) > 256 * 1024 * 1024:
                raise ValueError("expanded workbook exceeds preparation limit")
        workbook = openpyxl.load_workbook(
            io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
        )
        try:
            if source.sheet is None and len(workbook.sheetnames) != 1:
                raise ValueError("choose an explicit worksheet for a multi-sheet workbook")
            sheet = workbook[source.sheet or workbook.sheetnames[0]]
            rows = sheet.iter_rows()
            columns = [_text(cell.value) for cell in next(rows)]
            for index, row in enumerate(rows):
                if index >= MAX_ROWS or (index + 1) * len(columns) > MAX_CELLS:
                    raise ValueError("source exceeds preparation row/cell limit")
                values = []
                for column, cell in zip(columns, row, strict=True):
                    if cell.data_type in {"f", "e"}:
                        raise ValueError(
                            "workbook formulas/errors require a reviewed values-only export"
                        )
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        numeric[column] += 1
                    values.append(_text(cell.value))
                records.append(values)
        finally:
            workbook.close()
    else:
        raise ValueError("preparation supports CSV, TSV and explicitly selected XLSX sheets")
    if (
        not 1 <= len(columns) <= 256
        or any(empty(c) for c in columns)
        or len(columns) != len(set(columns))
    ):
        raise ValueError("source requires 1..256 distinct nonempty column names")
    if any(c.startswith("_rde_") for c in columns):
        raise ValueError("source column uses reserved preparation namespace _rde_")
    rows = [
        Row(
            i,
            dict(zip(columns, values, strict=True)),
            digest(dict(zip(columns, values, strict=True))),
        )
        for i, values in enumerate(records)
    ]
    return Table(source, columns, rows, dict(numeric))


def _key_stats(table, columns):
    values = [key(row.values, columns) for row in table.rows]
    counts = Counter(v for v in values if v is not None)
    return {
        "columns": columns,
        "null_rows": values.count(None),
        "distinct_keys": len(counts),
        "duplicate_rows": sum(n for n in counts.values() if n > 1),
        "unique": len(counts) == len(table.rows) and bool(table.rows),
    }


def source_schema(table, columns):
    """Describe the verified parser output, without inferring clinical roles."""
    source = table.source
    file_format = Path(source.path).suffix.lower().lstrip(".")
    schema = {
        "schema": "rde-source-table-schema-v1",
        "source_id": source.id,
        "source_sha256": source.sha256,
        "source_bytes": source.bytes,
        "rows": len(table.rows),
        "parser": {
            "format": file_format,
            "encoding": source.encoding if file_format != "xlsx" else None,
            "delimiter": source.delimiter if file_format != "xlsx" else None,
            "sheet": source.sheet if file_format == "xlsx" else None,
            "implicit_single_sheet": file_format == "xlsx" and source.sheet is None,
        },
        "storage_type": "text",
        "storage_representation": "All loaded cell values are strings. Delimited spelling is preserved; XLSX empty cells become empty strings, dates/times use ISO text, booleans true/false, and numeric cells use decimal text. This is storage normalization, not clinical classification or identifier validation.",
        "column_order": list(table.columns),
        "columns": [
            {
                "index": index,
                "name": column["name"],
                "storage_type": "text",
                "observed_constraints": {
                    "missing_rows": column["missing"],
                    "distinct_nonmissing": column["distinct"],
                    "numeric_excel_cells": column["numeric_excel_cells"],
                    "unique_nonmissing": column["distinct"] > 0
                    and column["distinct"] == len(table.rows) - column["missing"],
                    "no_missing": column["missing"] == 0,
                },
            }
            for index, column in enumerate(columns)
        ],
        "missing_convention": {
            "profile_predicate": "value is None, or a string equal to '' or satisfying Python str.isspace()",
            "loaded_null_cells": "empty strings",
            "whitespace": "counted as missing in profiling and key checks; stored spelling is preserved",
            "literal_tokens": "NA, N/A, NULL and other nonblank strings are retained as literal values; missing-token cleaning requires explicit preparation rules",
            "automatic_null_coercion": False,
        },
        "constraints_scope": "Observed file only; candidates do not declare primary/foreign keys or establish cross-system identity.",
        "clinical_semantics": "not_inferred",
    }
    return {**schema, "schema_sha256": digest(schema)}


def profile(table):
    columns = []
    for name in table.columns:
        values = [row.values[name] for row in table.rows]
        columns.append(
            {
                "name": name,
                "missing": sum(empty(v) for v in values),
                "distinct": len(set(v for v in values if not empty(v))),
                "numeric_excel_cells": table.numeric_cells.get(name, 0),
            }
        )
    ranked = sorted(
        table.columns,
        key=lambda c: (
            not bool(re.search(r"(^id$|_id$|code|編號|病歷|subject|patient|encounter)", c, re.I)),
            table.columns.index(c),
        ),
    )
    candidates = [_key_stats(table, [name]) for name in ranked[:16]]
    candidates += [_key_stats(table, list(pair)) for pair in itertools.combinations(ranked[:6], 2)]
    return {
        "id": table.source.id,
        "source_sha256": table.source.sha256,
        "rows": len(table.rows),
        "columns": columns,
        "source_schema": source_schema(table, columns),
        "candidate_keys": sorted(
            candidates,
            key=lambda c: (not c["unique"], c["null_rows"], c["duplicate_rows"], len(c["columns"])),
        )[:20],
        "identity_note": "Candidates are observed constraints, not proof that an identifier represents the same person across systems.",
    }


def inspect(tables):
    profiles = [profile(table) for table in tables.values()]
    relations = []
    for left, right in itertools.combinations(tables.values(), 2):
        shared = [c for c in left.columns if c in right.columns]
        candidates = [[c] for c in shared[:12]] + [
            list(c) for c in itertools.combinations(shared[:4], 2)
        ]
        for columns in candidates:
            lk = Counter(key(r.values, columns) for r in left.rows)
            rk = Counter(key(r.values, columns) for r in right.rows)
            lk.pop(None, None)
            rk.pop(None, None)
            common = lk.keys() & rk.keys()
            if not common:
                continue
            relations.append(
                {
                    "left_source": left.source.id,
                    "right_source": right.source.id,
                    "left_keys": columns,
                    "right_keys": columns,
                    "matched_left_rows": sum(lk[k] for k in common),
                    "matched_right_rows": sum(rk[k] for k in common),
                    "left_null_rows": sum(key(r.values, columns) is None for r in left.rows),
                    "right_null_rows": sum(key(r.values, columns) is None for r in right.rows),
                    "direct_join_rows": sum(lk[k] * rk[k] for k in common),
                    "left_unique": all(n == 1 for n in lk.values()),
                    "right_unique": all(n == 1 for n in rk.values()),
                }
            )
    return {
        "schema": "rde-preparation-inspection-v1",
        "sources": profiles,
        "candidate_relations": relations,
        "questions": [
            "What does one analysis row represent: a patient, encounter or event?",
            "Do the proposed identifiers refer to the same entity and namespace across systems?",
            "What is the index time, baseline/follow-up window and actual availability time?",
            "Does a missing linked record mean no observation, incomplete export or unknown coverage?",
        ],
        "capabilities": {
            "linkage": "deterministic composite keys; no probabilistic or fuzzy identity matching",
            "join_execution": "lookup or per-unit event aggregation; no Cartesian expansion",
            "formats": ["csv", "tsv", "xlsx"],
            "max_sources": 16,
            "max_rows_per_source": MAX_ROWS,
            "max_cells_per_source": MAX_CELLS,
        },
    }
