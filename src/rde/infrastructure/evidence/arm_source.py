"""Read immutable source coordinates without guessing headers, types or missingness."""

import csv
import hashlib
import posixpath
import warnings
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from rde.domain.models.evidence_arms import BinaryArmSpec

MAX_BYTES = 20 * 1024 * 1024
MAX_CELLS = 1_000_000


def _numeric_tokens(archive, sheet_name):
    """Keep numeric OOXML lexical precision instead of a binary-float round trip."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def xml(name):
        text = archive.read(name).decode("utf-8-sig")
        if "<!DOCTYPE" in text or "<!ENTITY" in text:
            raise ValueError("Source XML may not declare DTDs or entities.")
        return ET.fromstring(text)

    sheets = xml("xl/workbook.xml").findall(f"{{{main}}}sheets/{{{main}}}sheet")
    selected = [s for s in sheets if s.get("name") == sheet_name]
    if len(selected) != 1:
        raise ValueError("Select one exact worksheet explicitly.")
    links = [
        link
        for link in xml("xl/_rels/workbook.xml.rels")
        if link.get("Id") == selected[0].get(f"{{{rel}}}id")
    ]
    if (
        len(links) != 1
        or links[0].get("TargetMode") == "External"
        or links[0].get("Type") != rel + "/worksheet"
    ):
        raise ValueError("Selected sheet must have one internal worksheet relationship.")
    target = links[0].get("Target", "")
    part = posixpath.normpath(posixpath.join("/xl", target)).lstrip("/")
    cells = xml(part).findall(f"{{{main}}}sheetData/{{{main}}}row/{{{main}}}c")
    seen, numbers = set(), {}
    for cell in cells:
        coordinate = cell.get("r")
        if not coordinate or coordinate in seen:
            raise ValueError("Source worksheet cells need unique explicit coordinates.")
        seen.add(coordinate)
        value = cell.find(f"{{{main}}}v")
        if (
            cell.get("t", "n") == "n"
            and cell.find(f"{{{main}}}f") is None
            and value is not None
            and value.text is not None
        ):
            numbers[coordinate] = value.text
    return part, numbers


def source_grid(data: bytes, filename: str, spec: BinaryArmSpec):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Source must be nonempty and no larger than 20 MiB.")
    if Path(filename).name != filename or "\\" in filename or len(filename) > 200:
        raise ValueError("Use a plain source filename.")
    suffix = Path(filename).suffix.lower()
    grids, formulas, merges, ends, parser_warnings = [], [], [], [], []
    worksheet_part = None
    if suffix in {".csv", ".tsv"}:
        if spec.sheet is not None:
            raise ValueError("Delimited sources do not have worksheets.")
        text = data.decode("utf-8-sig", errors="strict")
        if "\x00" in text:
            raise ValueError("NUL is not permitted in a delimited source.")
        reader = csv.reader(
            StringIO(text, newline=""), delimiter="\t" if suffix == ".tsv" else ",", strict=True
        )
        for row in reader:
            grids.append(row)
            ends.append(reader.line_num)
            if len(grids) > 10000 or len(row) > 100:
                raise ValueError("Source exceeds 10000 logical rows or 100 columns.")
        width = max(map(len, grids), default=0)
        if any(len(row) not in {0, width} for row in grids):
            raise ValueError(
                "Delimited records must have equal widths; blank records retain their positions."
            )
        grids = [row if row else [""] * width for row in grids]
    elif suffix == ".xlsx":
        import openpyxl

        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 64 * 1024 * 1024:
                raise ValueError("XLSX expansion exceeds the source budget.")
            if any(e.flag_bits & 1 for e in entries):
                raise ValueError("Encrypted XLSX entries are not supported.")
            if len({e.filename for e in entries}) != len(entries):
                raise ValueError("Duplicate XLSX archive members are ambiguous.")
            worksheet_part, numbers = _numeric_tokens(archive, spec.sheet)
        with warnings.catch_warnings(record=True) as observed:
            warnings.simplefilter("always")
            workbook = openpyxl.load_workbook(BytesIO(data), data_only=False, keep_links=False)
        parser_warnings = [
            {"category": warning.category.__name__, "message": str(warning.message)}
            for warning in observed
        ]
        try:
            if spec.sheet is None or spec.sheet not in workbook.sheetnames:
                raise ValueError("Select the exact worksheet explicitly.")
            sheet = workbook[spec.sheet]
            if (
                sheet.max_row > 10000
                or sheet.max_column > 100
                or sheet.max_row * sheet.max_column > MAX_CELLS
            ):
                raise ValueError("Worksheet exceeds the physical source budget.")
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        formulas.append({"row": cell.row, "column": cell.column})
                    if value is None:
                        values.append("")
                    elif isinstance(value, (date, datetime)):
                        values.append(value.isoformat())
                    elif isinstance(value, bool):
                        values.append("TRUE" if value else "FALSE")
                    elif cell.coordinate in numbers:
                        values.append(numbers[cell.coordinate])
                    else:
                        values.append(str(value))
                grids.append(values)
            merges = [
                {
                    "first_row": r.min_row,
                    "last_row": r.max_row,
                    "first_column": r.min_col,
                    "last_column": r.max_col,
                }
                for r in sheet.merged_cells.ranges
            ]
        finally:
            workbook.close()
    else:
        raise ValueError("Arm preparation accepts UTF-8 CSV/TSV or XLSX only.")
    if not grids or spec.last_data_row > len(grids):
        raise ValueError("Selected data rows are not all present in the source.")
    if sum(map(len, grids)) > MAX_CELLS or any(len(v) > 8000 for row in grids for v in row):
        raise ValueError("Source text exceeds the bounded review budget.")
    return {
        "schema": "evidence-arm-grid-v1",
        "source_sha256": hashlib.sha256(data).hexdigest(),
        "filename": filename,
        "sheet": spec.sheet,
        "worksheet_part": worksheet_part,
        "numeric_representation": "XLSX numeric v tokens are preserved verbatim; dates remain ISO values; no implicit cached formulas. CSV/TSV retains field strings.",
        "parser_warnings": parser_warnings,
        "row_count": len(grids),
        "column_count": max(map(len, grids)),
        "rows": grids,
        "formula_cells": formulas,
        "merged_ranges": merges,
        "csv_record_end_lines": ends,
        "coordinates": "one-based worksheet rows/columns; CSV rows are logical records, not physical lines",
    }
