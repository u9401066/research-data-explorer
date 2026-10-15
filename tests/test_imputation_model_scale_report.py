"""A binary predictor does not turn a continuous-outcome MI model into logistic."""

from dataclasses import replace
from pathlib import Path
import re

from docx import Document
import pytest

from rde.application.pipeline import PipelinePhase
from rde.interface.mcp.tools.clinical_tools import clinical_records, verify_clinical_artifacts
from test_clinical_formal_exports import call
from test_clinical_imputation import fixture
from test_clinical_survival import clinical_project


@pytest.mark.parametrize("distribution", ["gaussian", "binomial"])
def test_actual_mi_docx_and_publication_respect_outcome_model_scale(
    tmp_path, monkeypatch, distribution
):
    from rde.infrastructure.clinical import imputation

    frame, spec = fixture(distribution)
    # Chinese original units belong in source context, not untranslated English captions.
    spec = replace(spec, outcome_unit="合成結果單位", context="工程用模型尺度驗證，非臨床結論")
    project, store, dataset, spec = clinical_project(tmp_path, frame=frame, options=spec)
    call("run_clinical_study", {"dataset_id": dataset.id, "clinical_options": spec.to_dict()})
    record = clinical_records(store)[0]
    frozen = store.get_path(PipelinePhase.EXECUTE_EXPLORATION, record["artifact"]).read_bytes()
    monkeypatch.setattr(
        imputation, "run_imputation", lambda *a, **k: pytest.fail("Export refitted MI")
    )
    call("collect_results", {"project_id": project.id})
    call("assemble_report", {"project_id": project.id})
    call("export_report", {"project_id": project.id, "formats": "docx"})
    exports = project.artifacts_dir / "exports"
    formal = (exports / "eda_report_formal_source.md").read_text()
    document = Document(exports / "eda_report.docx")
    docx_text = "\n".join(p.text for p in document.paragraphs)
    coefficient_tables = [
        table
        for table in document.tables
        if any("within variance" in cell.text for cell in table.rows[0].cells)
    ]
    assert len(coefficient_tables) == 1
    assert [row.cells[0].text for row in coefficient_tables[0].rows[1:]] == ["intercept", "x", "b"]
    effect = next(f for f in record["figures"] if f["plot_type"] == "clinical_imputation_effects_1")
    pub = effect["publication"]
    caption = pub["caption_en"]
    assert re.search(r"[\u3400-\u9fff]", caption) is None
    assert pub["source_context"]["outcome_unit"] == "合成結果單位"
    assert caption in Path(pub["files"]["caption"]).read_text()
    for text in [formal, docx_text]:
        assert "不是邊際平均差或因果效果" in text or "不是邊際 OR、風險比或風險差" in text
        assert "Binary 係數是 log-odds" not in text
        assert "Binary values are log-odds coefficients" not in text
        if distribution == "gaussian":
            assert "合成結果單位" in text
            assert "二元預測變項" in text and "1 相對 0" in text
            assert "條件平均差" in text and "連續預測變項每增加一個來源單位" in text
        else:
            assert "條件 log-odds" in text and "條件 OR" in text
    if distribution == "gaussian":
        assert "original outcome units" in caption
        assert "continuous predictor per one source unit" in caption
        assert "binary predictor coded 1 versus 0" in caption
        assert "conditional mean differences" in caption
        assert "not marginal effects" in caption
        assert "log-odds" not in caption
    else:
        assert "conditional log-odds" in caption
        assert "not marginal odds ratios" in caption
        assert "reference log-odds" in caption
        assert "conditional mean differences" not in caption
    harness = store.load(PipelinePhase.REPORT_ASSEMBLY, "figure_interpretation_harness.json")
    entry = next(e for e in harness["entries"] if e["plot_type"] == effect["plot_type"])
    assert "binary predictor does not determine the model link" in entry["validity_caveat"]
    assert verify_clinical_artifacts(record, project.output_dir)
    assert (
        store.get_path(PipelinePhase.EXECUTE_EXPLORATION, record["artifact"]).read_bytes() == frozen
    )
