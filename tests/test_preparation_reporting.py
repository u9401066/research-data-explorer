"""Report preparation claims require the exact retained release, not filenames."""

import copy
import hashlib
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest

from rde.application.pipeline import PipelinePhase
from rde.infrastructure.adapters.descriptive_study import markdown, summarize
from rde.infrastructure.adapters.docx_exporter import DocxExporter
from rde.infrastructure.adapters.markdown_renderer import MarkdownReportRenderer
from rde.infrastructure.persistence.artifact_store import ArtifactStore
from rde.infrastructure.preparation import portable
from rde.infrastructure.preparation.reporting import preparation_evidence, preparation_markdown
from rde.interface.mcp.tools.report_tools import (
    _build_formal_research_report,
    _build_claim_provenance_manifest,
    _formal_abstract,
    _formal_data_quality,
    _format_analyses,
)
from test_preparation_intake import released


def imported(tmp_path):
    _, plan, release, bundle = released(tmp_path)
    project = SimpleNamespace(id="analysis", output_dir=tmp_path / "analysis", dataset_ids=["d"])
    project.artifacts_dir = project.output_dir / "artifacts"
    binding = portable.import_bundle(
        project, bundle["path"], bundle["sha256"], release["receipt_sha256"]
    )
    intake = dict(
        preparation=binding,
        row_count=release["population"]["rows"],
        source={k: release["dataset"][k] for k in ("bytes", "sha256")},
        dataset_id="d",
    )
    return project, intake, plan


def test_exact_retained_preparation_methods_denominators_and_docx(tmp_path):
    from docx import Document

    project, intake, plan = imported(tmp_path)
    folder = Path(intake["preparation"]["folder"])
    original = {
        str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.rglob("*") if p.is_file()
    }
    evidence = preparation_evidence(project, intake)
    content = preparation_markdown(evidence)
    assert "6 個分析單位、6 個受試者識別" in content
    assert "synthetic_only" in content
    assert "synthetic_oracle" in content
    assert "matched_partial 表示已配對" in content
    assert "事件配對" in content and "有紀錄受試者" in content
    assert "outside_window" in content and "not_exported" in content
    assert "lab_mean" in content and "mean" in content
    assert plan["receipt_sha256"] in content
    for item in plan["source_snapshots"]:
        assert item["source"]["sha256"] in content
        assert item["source"]["path"] not in content

    store = ArtifactStore(project.artifacts_dir)
    store.save(PipelinePhase.DATA_INTAKE, "intake_report.json", intake)
    result = summarize(
        pd.DataFrame({"value": [0.9, 1.3, 2.4000000000000004]}),
        dict(columns=["value"], identifiers=[], unresolved=[]),
    )
    store.save(
        PipelinePhase.COLLECT_RESULTS,
        "results_summary.json",
        {
            "descriptive_studies": [
                {
                    "result": result,
                    "receipt_sha256": "f" * 64,
                    "summary_markdown": "obsolete presentation",
                }
            ]
        },
    )
    report = _build_formal_research_report(project, store, "# Prepared report")
    rendered = MarkdownReportRenderer().render_markdown(report)
    assert "已發布、可追溯" in rendered
    assert "原始資料列的描述性研究" not in rendered
    assert "obsolete presentation" not in rendered
    assert "q3=2.4000000000000004" not in rendered
    assert plan["receipt_sha256"] in rendered
    manifest = _build_claim_provenance_manifest(project, store, report_source="test")
    claim = next(c for c in manifest["claims"] if c["claim_type"] == "prepared_data_methods")
    assert any(e.get("receipt_sha256") == plan["receipt_sha256"] for e in claim["evidence"])
    assert all(not e["artifact"].startswith("/") for e in claim["evidence"])
    path = tmp_path / "prepared.docx"
    DocxExporter().export_docx(report, path)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    assert "lab_mean" in text and plan["receipt_sha256"] in text
    assert "準備階段納排分母" in text
    assert original == {
        str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.rglob("*") if p.is_file()
    }


def test_automated_approval_is_not_promoted_to_clinical_validation(tmp_path):
    from test_preparation_automated_review import prepared, call, approval_request
    from rde.infrastructure.preparation import workflow

    source_project, plan, _ = prepared(tmp_path)
    audit = call(source_project, plan, "automated_review")
    approval = workflow.execute(source_project, approval_request(plan, audit))
    release = call(source_project, plan, "release", approval_sha256=approval["receipt_sha256"])
    bundle = call(source_project, plan, "bundle", release_sha256=release["receipt_sha256"])
    project = SimpleNamespace(output_dir=tmp_path / "analysis")
    binding = portable.import_bundle(
        project, bundle["path"], bundle["sha256"], release["receipt_sha256"]
    )
    evidence = preparation_evidence(
        project, dict(preparation=binding, source=release["dataset"], row_count=6)
    )
    content = preparation_markdown(evidence)
    assert "核准種類：automated_checks" in content
    assert "clinical_validation：not_performed" in content
    assert "不能解讀為人工臨床核准" in content
    (Path(binding["folder"]) / "automated_review.json").unlink()
    with pytest.raises(ValueError, match="PREPARATION_REPORT_EVIDENCE_INVALID"):
        preparation_evidence(
            project, dict(preparation=binding, source=release["dataset"], row_count=6)
        )


def test_pdf_html_uses_real_preparation_tables_and_headings(tmp_path):
    from rde.domain.models.report import EDAReport, ReportSection

    report = EDAReport(
        id="report", dataset_id="dataset", project_id="project", title="Prepared flow"
    )
    report.add_section(
        ReportSection(
            section_id="methods",
            title="Methods",
            content="#### Flow\n\n| Stage | Units |\n| --- | ---: |\n| linked | 275 |\n| window | 35 |",
            order=1,
        )
    )
    html = DocxExporter()._build_pdf_html(report)
    assert "<h4>Flow</h4>" in html
    assert "<th>Stage</th>" in html and "<td>275</td>" in html
    assert "<p>|" not in html
    assert html.count("<table>") == 1
    assert "<thead><tr>" in html and "</tbody></table>" in html
    assert "page-break-inside: avoid" in html


@pytest.mark.parametrize(
    "fault",
    [
        "missing_bundle",
        "tampered_trace",
        "tampered_spec",
        "wrong_source",
        "wrong_population",
        "sibling_folder",
        "missing_source_identity",
    ],
)
def test_prepared_report_fails_closed(tmp_path, fault):
    project, intake, _ = imported(tmp_path)
    folder = Path(intake["preparation"]["folder"])
    if fault == "missing_bundle":
        (folder / "release.json").unlink()
    elif fault == "tampered_trace":
        (folder / "trace.jsonl").write_text("{}\n")
    elif fault == "tampered_spec":
        (folder / "specification.json").write_text("{}")
    elif fault == "wrong_source":
        intake["source"]["sha256"] = "0" * 64
    elif fault == "wrong_population":
        intake["preparation"]["population"]["subjects"] = 999
    elif fault == "sibling_folder":
        project.output_dir = tmp_path / "other-project"
    else:
        del intake["source"]
    with pytest.raises(ValueError, match="PREPARATION_REPORT_EVIDENCE_INVALID"):
        preparation_evidence(project, intake)
    store = ArtifactStore(project.artifacts_dir)
    store.save(PipelinePhase.DATA_INTAKE, "intake_report.json", intake)
    with pytest.raises(ValueError, match="PREPARATION_REPORT_EVIDENCE_INVALID"):
        _build_formal_research_report(project, store, "# Must not render verified methods")


def test_ordinary_input_and_unperformed_checks_do_not_gain_prepared_or_clinical_claims():
    assert preparation_evidence(SimpleNamespace(), {}) is None
    abstract = _formal_abstract(None, {}, {}, {"descriptive_studies": [True]})
    assert "已發布" not in abstract and "原始資料列" not in abstract
    quality = _formal_data_quality(
        {},
        {
            "checks": [
                {"name": "sparse groups", "passed": False, "applicable": False},
                {"name": "collinearity", "passed": False, "status": "skipped"},
                {"name": "source integrity", "passed": False},
            ]
        },
    )
    assert "sparse group、缺失模式與共線性仍限制正式推論" not in quality
    assert "sparse groups 未通過" not in quality
    assert "collinearity 未通過" not in quality
    assert "source integrity 未通過" in quality
    assert "未執行／不適用" in quality


def test_numeric_display_does_not_change_receipt_values_or_saved_summary():
    result = summarize(
        pd.DataFrame({"v": [0.4, 0.9, 1.3, 2.4000000000000004, 15.0]}),
        dict(columns=["v"], identifiers=[], unresolved=[]),
    )
    original = copy.deepcopy(result)
    assert result["columns"][0]["summary"]["q3"] == 2.4000000000000004
    old = dict(result=result, receipt_sha256="f" * 64, summary_markdown="q3=2.4000000000000004")
    rendered = _format_analyses({"descriptive_studies": [old]})
    assert "q3=2.4" in rendered and "q3=2.4000000000000004" not in rendered
    assert "本描述步驟沒有刪除或合併資料" in markdown(result)
    assert result == original
    assert old["summary_markdown"] == "q3=2.4000000000000004"
