"""Source/approval/restart edges through MCP; independent R checks run via smoke."""

import asyncio
import csv
import json
import uuid
from zipfile import ZipFile
from copy import deepcopy
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl
import pytest

from rde.application.session import get_session
from rde.domain.models.evidence_arms import BinaryArmSpec
from rde.infrastructure.evidence import arm_preparation as engine
from rde.infrastructure.evidence import arm_source
from rde.infrastructure.evidence import arm_workflow as workflow
from rde.interface.mcp.server import create_server

FIXTURES = Path(__file__).parent / "fixtures/evidence-arms"


def spec():
    return json.loads((FIXTURES / "spec.json").read_text())


def call(arguments, error=False):
    response = asyncio.run(create_server().call_tool("evidence_arm_preparation", arguments))
    assert response.is_error is error, response.content
    return response.content[0].text if error else json.loads(response.content[0].text)


def setup(options=None):
    asyncio.run(
        create_server().call_tool(
            "init_project",
            {"name": "Synthetic arms QA", "research_question": "Software verification only"},
        )
    )
    project = get_session().get_project()
    identity = {"project_id": project.id}
    prep = str(uuid.uuid4())
    incoming = project.output_dir / "incoming/evidence-arms" / prep / "source.csv"
    incoming.parent.mkdir(parents=True)
    incoming.write_bytes((FIXTURES / "source.csv").read_bytes())
    request = {
        "op": "draft",
        "preparation_id": prep,
        "filename": incoming.name,
        "source_sha256": workflow.w.file_hash(incoming),
        "specification": options or spec(),
    }
    plan = call({**identity, "request": request})
    return project, identity, request, plan


def approve(identity, request, plan, error=False):
    return call(
        {
            **identity,
            "request": {
                "op": "approve",
                "preparation_id": request["preparation_id"],
                "expected_plan_sha256": plan["receipt_sha256"],
                "review": {
                    "reviewer": "Authorized engineering QA",
                    "note": "Synthetic assumptions reviewed; not clinical approval.",
                    "confirmations": {name: True for name in plan["required_confirmations"]},
                },
            },
        },
        error=error,
    )


def execute_request(request, plan, approval):
    return {
        "op": "execute",
        "preparation_id": request["preparation_id"],
        "expected_plan_sha256": plan["receipt_sha256"],
        "expected_approval_sha256": approval["receipt_sha256"],
        "run_id": str(uuid.uuid4()),
    }


def test_mcp_review_approval_run_paging_and_immutable_replay(monkeypatch):
    project, identity, request, plan = setup()
    assert plan["ready_for_approval"] and plan["summary"]["planned_contrasts"] == 4
    assert plan["summary"]["source_rows"] == 11 and plan["summary"]["excluded_studies"] == 2
    assert call({**identity, "request": request}) == plan
    denied = execute_request(request, plan, {"receipt_sha256": "0" * 64})
    call({**identity, "request": denied}, error=True)
    approval = approve(identity, request, plan)
    assert approve(identity, request, plan) == approval
    run_request = execute_request(request, plan, approval)
    run = call({**identity, "request": run_request})
    full = ""
    page_request = {
        "op": "read",
        "preparation_id": request["preparation_id"],
        "part": "result",
        "run_id": run["run_id"],
        "text_limit": 701,
    }
    while True:
        page = call({**identity, "request": page_request})
        full += page["text_excerpt"]
        if page["next_text_offset"] is None:
            break
        page_request.update(
            text_offset=page["next_text_offset"], expected_text_sha256=page["text_sha256"]
        )
    assert workflow.w.sha(full.encode()) == page["text_sha256"]
    result = json.loads(full)
    assert len(result["contrasts"]) == 4
    assert result["derivations"][0]["arms"][0]["events"] == "5.5"
    assert result["derivations"][0]["arms"][0]["source_rows"] == [3, 4]
    assert result["derivations"][0]["arms"][0]["corrected_events"] == "6.0"
    assert all(a["corrected_total"] != a["total"] for a in result["derivations"][0]["arms"])
    # Shared-arm algebra: the three paired effects close exactly within rounding.
    ab, ac, bc = result["contrasts"][:3]
    assert ab["effect"] + bc["effect"] == pytest.approx(ac["effect"], abs=1e-14)
    monkeypatch.setattr(
        engine, "calculate", lambda *_: pytest.fail("Completed replay must not calculate")
    )
    assert call({**identity, "request": run_request}) == run
    assert project.dataset_ids == []
    page_request["expected_text_sha256"] = "0" * 64
    call({**identity, "request": page_request}, error=True)
    (project.output_dir / run["contrast_path"]).write_text("tampered\n")
    assert "arm evidence file size changed" in call(
        {**identity, "request": run_request}, error=True
    )


@pytest.mark.parametrize(
    "field,value,code",
    [
        ("merge_policy", "reject", "merge_not_approved"),
        ("zero_cells", "reject", "zero_cells"),
        ("fractional_events", "reject", "invalid_counts"),
        ("uninformative_studies", "reject", "uninformative_study"),
        ("insufficient_treatments", "reject", "insufficient_treatments"),
    ],
)
def test_unresolved_policy_cannot_be_approved(field, value, code):
    options = spec()
    options[field] = value
    project, identity, request, plan = setup(options)
    assert not plan["ready_for_approval"]
    reviewed = workflow.w.load_json(
        workflow.directory(project, request["preparation_id"]) / "review.json"
    )
    assert code in {e["code"] for e in reviewed["errors"]}
    assert "resolve source review errors" in approve(identity, request, plan, error=True)


@pytest.mark.parametrize(
    "field,value",
    [
        ("independent_parallel_arms", 1),
        ("independent_parallel_arms", False),
        ("outcome", " "),
        ("unmentioned_default", True),
        (
            "missing_codes",
            [{"value": "", "meaning": "blank"}, {"value": "0", "meaning": "missing"}],
        ),
    ],
)
def test_strict_spec_does_not_coerce_or_silently_default(field, value):
    options = spec()
    options[field] = value
    with pytest.raises(ValueError):
        BinaryArmSpec.model_validate(options)


@pytest.mark.parametrize(
    "mutation,code",
    [
        ("duplicate", "duplicate_arm"),
        ("denominator", "invalid_counts"),
        ("flag", "unknown_imputation"),
        ("fraction_without_flag", "invalid_counts"),
        ("unknown_treatment", "unknown_treatment"),
        ("metadata", "inconsistent_metadata"),
    ],
)
def test_raw_arm_edges_retain_source_and_stop(mutation, code):
    options = BinaryArmSpec.model_validate(spec())
    grid = arm_source.source_grid((FIXTURES / "source.csv").read_bytes(), "source.csv", options)
    if mutation == "duplicate":
        grid["rows"][3][1] = "A.low"
    elif mutation == "denominator":
        grid["rows"][2][4] = "2"
    elif mutation == "flag":
        grid["rows"][2][5] = ""
    elif mutation == "fraction_without_flag":
        grid["rows"][2][5] = "0"
    elif mutation == "unknown_treatment":
        grid["rows"][2][2] = "unreviewed"
    else:
        grid["rows"][2][6] = "another report"
    reviewed = engine.review(grid, options)
    assert not reviewed["ready_for_approval"] and code in {e["code"] for e in reviewed["errors"]}
    assert reviewed["arms"][4]["missing"]["events"] == "Not available in source."
    with pytest.raises(ValueError, match="Unresolved"):
        engine.calculate(reviewed, options)


def test_xlsx_multilevel_headers_exact_cells_formula_and_merge_guards():
    options = spec()
    options["sheet"] = "DATA"
    parsed = BinaryArmSpec.model_validate(options)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "DATA"
    for row in csv.reader(StringIO((FIXTURES / "source.csv").read_text())):
        sheet.append(row)
    sheet.merge_cells("A1:G1")
    data = BytesIO()
    book.save(data)
    grid = arm_source.source_grid(data.getvalue(), "source.xlsx", parsed)
    assert engine.review(grid, parsed)["ready_for_approval"]
    assert grid["rows"][2][3] == "2.5"
    sheet["D3"] = "=2.5"
    sheet.merge_cells("F8:F9")
    data = BytesIO()
    book.save(data)
    reviewed = engine.review(arm_source.source_grid(data.getvalue(), "source.xlsx", parsed), parsed)
    assert {"formula_cell", "merged_cell"} <= {e["code"] for e in reviewed["errors"]}


def test_csv_blank_records_and_embedded_newlines_keep_logical_coordinates():
    options = spec()
    options.update(first_data_row=1, last_data_row=3)
    grid = arm_source.source_grid(
        b'"first\nline",x\n\nlast,z\n', "source.csv", BinaryArmSpec.model_validate(options)
    )
    assert grid["rows"] == [["first\nline", "x"], ["", ""], ["last", "z"]]
    assert grid["csv_record_end_lines"] == [2, 3, 4]


def test_xlsx_numeric_lexical_precision_survives_without_float_rounding():
    options = spec()
    options.update(sheet="DATA", first_data_row=1, last_data_row=1)
    book = openpyxl.Workbook()
    book.active.title = "DATA"
    book.active["A1"] = 1.25
    original, rewritten = BytesIO(), BytesIO()
    book.save(original)
    precise = "1.23456789012345678901"
    with ZipFile(original) as source, ZipFile(rewritten, "w") as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v>1.25</v>", ("<v>" + precise + "</v>").encode())
            target.writestr(member, data)
    grid = arm_source.source_grid(
        rewritten.getvalue(), "source.xlsx", BinaryArmSpec.model_validate(options)
    )
    assert grid["rows"][0][0] == precise
    assert grid["worksheet_part"] == "xl/worksheets/sheet1.xml"


def test_unadjudicated_real_source_vocabulary_is_a_draft_not_an_exclusion():
    options = spec()
    for treatment in options["treatments"]:
        treatment.update(decision="unresolved", reason="Source eligibility not yet adjudicated.")
    parsed = BinaryArmSpec.model_validate(options)
    reviewed = engine.review(
        arm_source.source_grid((FIXTURES / "source.csv").read_bytes(), "source.csv", parsed), parsed
    )
    assert not reviewed["ready_for_approval"]
    assert reviewed["summary"]["included_studies"] == 0
    assert reviewed["summary"]["excluded_studies"] == 0
    assert reviewed["summary"]["arm_decisions"] == {"unresolved": 11}
    assert all(study["decision"] == "unresolved" for study in reviewed["studies"])
    assert sum(error["code"] == "unresolved_eligibility" for error in reviewed["errors"]) == 11


@pytest.mark.parametrize(
    "problem", ["source", "symlink", "changed_implementation", "partial_run", "failed_run"]
)
def test_immutable_source_and_failed_attempt_guards(problem, monkeypatch):
    project, identity, request, plan = setup()
    approval = approve(identity, request, plan)
    run_request = execute_request(request, plan, approval)
    root = workflow.directory(project, request["preparation_id"])
    if problem == "source":
        (root / "source/source.csv").write_text("changed")
    elif problem == "symlink":
        (root / "source/source.csv").rename(root / "original.csv")
        (root / "source/source.csv").symlink_to(root / "original.csv")
    elif problem == "changed_implementation":
        monkeypatch.setattr(workflow, "implementation", dict)
    elif problem == "partial_run":
        (root / "runs" / run_request["run_id"]).mkdir(parents=True)
    else:

        def fail(*_):
            raise ValueError("Injected engineering failure")

        monkeypatch.setattr(engine, "calculate", fail)
    call({**identity, "request": run_request}, error=True)
    if problem == "failed_run":
        failure = root / "runs" / run_request["run_id"] / "failure.json"
        assert failure.is_file()
        original = failure.read_bytes()
        assert "failed or interrupted" in call({**identity, "request": run_request}, error=True)
        assert failure.read_bytes() == original


def test_contract_and_request_reject_path_and_unknown_parameters():
    contract_result = call({"project_id": "", "request": {"op": "contract"}})
    assert contract_result["request_schema"]["discriminator"]["propertyName"] == "op"
    _, identity, request, _ = setup()
    changed = deepcopy(request)
    changed["filename"] = "../source.csv"
    call({**identity, "request": changed}, error=True)
    request["unexpected"] = True
    call({**identity, "request": request}, error=True)
